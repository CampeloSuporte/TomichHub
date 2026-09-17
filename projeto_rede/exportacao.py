"""
Exportação do documento de rede: HTML de impressão, PDF e DOCX.

- PDF: o mesmo HTML da pré-visualização impresso pelo Google Chrome headless
  (`--print-to-pdf`), com cabeçalho/rodapé via `@page` — sem dependência
  nova no servidor e com fidelidade ao que o usuário vê.
- DOCX: conversão do HTML sanitizado (subconjunto conhecido, ver
  `sanitizar.py`) para python-docx, com capa, sumário e estilos próprios.
"""
import io
import os
import re
import shutil
import subprocess
import tempfile
from html.parser import HTMLParser

from django.template.loader import render_to_string
from django.utils import timezone

from .sanitizar import limpar_html

COR_PRIMARIA = '0F6E7A'
COR_ZEBRA = 'F1F6F7'
COR_CALLOUT = 'E3EFF2'
COR_RISCO = 'FDF1E3'
CORES_SEV = {
    'sev-critica': 'B42318', 'sev-alta': 'C4320A', 'sev-media': 'B54708',
    'sev-baixa': '067647', 'sev-info': '475467',
}
CHROME = ['/usr/bin/google-chrome', '/usr/bin/google-chrome-stable', '/usr/bin/chromium', '/usr/bin/chromium-browser']


class ErroExportacao(Exception):
    pass


# ═══════════════════════════════════════════════════════════════════════════
# Numeração
# ═══════════════════════════════════════════════════════════════════════════

def secoes_numeradas(secoes):
    """[{numero, titulo, html}] com h3 prefixados por 'n.m'. Anexos não
    entram na numeração principal."""
    saida = []
    n = 0
    for s in secoes:
        titulo = s.get('titulo') or ''
        anexo = titulo.lower().startswith('anexo')
        numero = ''
        if not anexo:
            n += 1
            numero = str(n)
        contador = {'i': 0}

        def _num(m, numero=numero, contador=contador):
            contador['i'] += 1
            if not numero:
                return m.group(0)
            return f'{m.group(1)}<span class="num">{numero}.{contador["i"]}</span> '

        html = re.sub(r'(<h3[^>]*>)', _num, limpar_html(s.get('html') or ''))
        saida.append({'id': s.get('id'), 'numero': numero, 'titulo': titulo, 'html': html, 'anexo': anexo})
    return saida


def _campos(doc):
    from .documentos import campos
    return campos(doc.tipo)


def _css_str(texto):
    """Conteúdo seguro para uma string CSS entre aspas duplas."""
    saida = []
    for ch in ' '.join(str(texto or '').split()):
        if ch in '"\\' or ch in '<>&' or ord(ch) < 32:
            saida.append(f'\\{ord(ch):x} ')
        else:
            saida.append(ch)
    return '"' + ''.join(saida) + '"'


def contexto_documento(doc):
    meta = doc.metadados or {}
    return {
        'doc': doc,
        'meta': meta,
        'css_topo': _css_str(f'{meta.get("empresa", "")} · {meta.get("documento") or doc.titulo}'),
        'css_rodape': _css_str(f'{meta.get("classificacao", "")} · v{doc.versao}'),
        'campos': [(rotulo, meta.get(chave, '')) for chave, rotulo in _campos(doc)],
        'principio_titulo': meta.get('principio_titulo') or 'Princípio do documento.',
        'secoes': secoes_numeradas(doc.secoes or []),
        'gerado_em': timezone.localtime(),
    }


def html_documento(doc, para_pdf=False):
    ctx = contexto_documento(doc)
    ctx['para_pdf'] = para_pdf
    return render_to_string('projeto_rede/documento_impressao.html', ctx)


def nome_arquivo(doc, ext):
    base = f'{doc.titulo}-v{doc.versao}'
    base = re.sub(r'[^\w.\-]+', '_', base, flags=re.ASCII).strip('_') or 'documento'
    return f'{base}.{ext}'


# ═══════════════════════════════════════════════════════════════════════════
# PDF
# ═══════════════════════════════════════════════════════════════════════════

def gerar_pdf(doc):
    binario = next((c for c in CHROME if os.path.exists(c)), None) or shutil.which('google-chrome')
    if not binario:
        raise ErroExportacao('Google Chrome não encontrado no servidor para gerar o PDF.')
    html = html_documento(doc, para_pdf=True)
    with tempfile.TemporaryDirectory(prefix='asis_pdf_') as tmp:
        entrada = os.path.join(tmp, 'documento.html')
        saida = os.path.join(tmp, 'documento.pdf')
        with open(entrada, 'w', encoding='utf-8') as fh:
            fh.write(html)
        cmd = [
            binario, '--headless=new', '--no-sandbox', '--disable-gpu', '--disable-extensions',
            '--no-first-run', '--no-default-browser-check', '--disable-dev-shm-usage',
            '--no-pdf-header-footer', '--run-all-compositor-stages-before-draw',
            f'--user-data-dir={os.path.join(tmp, "perfil")}',
            f'--print-to-pdf={saida}', f'file://{entrada}',
        ]
        env = dict(os.environ, HOME=tmp, XDG_CONFIG_HOME=tmp, XDG_CACHE_HOME=tmp)
        try:
            proc = subprocess.run(cmd, capture_output=True, timeout=90, env=env)
        except subprocess.TimeoutExpired:
            raise ErroExportacao('O Chrome não terminou de gerar o PDF em 90 s.')
        if not os.path.exists(saida) or os.path.getsize(saida) == 0:
            detalhe = (proc.stderr or b'').decode('utf-8', 'replace').strip().splitlines()[-1:] or ['']
            raise ErroExportacao(f'Falha ao gerar o PDF: {detalhe[0][:200]}')
        with open(saida, 'rb') as fh:
            return fh.read()


# ═══════════════════════════════════════════════════════════════════════════
# DOCX
# ═══════════════════════════════════════════════════════════════════════════

class _No:
    __slots__ = ('tag', 'attrs', 'filhos')

    def __init__(self, tag, attrs=None):
        self.tag = tag
        self.attrs = dict(attrs or {})
        self.filhos = []

    @property
    def classes(self):
        return set((self.attrs.get('class') or '').split())


class _Arvore(HTMLParser):
    VAZIAS = {'br'}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.raiz = _No('raiz')
        self.pilha = [self.raiz]

    def handle_starttag(self, tag, attrs):
        no = _No(tag, attrs)
        self.pilha[-1].filhos.append(no)
        if tag not in self.VAZIAS:
            self.pilha.append(no)

    def handle_startendtag(self, tag, attrs):
        self.pilha[-1].filhos.append(_No(tag, attrs))

    def handle_endtag(self, tag):
        for i in range(len(self.pilha) - 1, 0, -1):
            if self.pilha[i].tag == tag:
                del self.pilha[i:]
                break

    def handle_data(self, data):
        self.pilha[-1].filhos.append(data)


def _arvore(html):
    p = _Arvore()
    p.feed(html or '')
    p.close()
    return p.raiz


def _texto(no):
    if isinstance(no, str):
        return no
    if no.tag == 'br':
        return '\n'
    return ''.join(_texto(f) for f in no.filhos)


def _sombrear(celula, cor):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), cor)
    tc_pr.append(shd)


def _borda_esquerda(celula, cor, espessura=24):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = celula._tc.get_or_add_tcPr()
    bordas = OxmlElement('w:tcBorders')
    for lado in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{lado}')
        if lado == 'left':
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), str(espessura))
            el.set(qn('w:color'), cor)
        else:
            el.set(qn('w:val'), 'nil')
        bordas.append(el)
    tc_pr.append(bordas)


def _repetir_cabecalho(linha):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tr_pr = linha._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    tr_pr.append(el)


def _campo(run, instrucao):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for tipo, texto in (('begin', None), (None, instrucao), ('separate', None), (None, '1'), ('end', None)):
        if tipo:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tipo)
        elif texto == instrucao:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = f' {instrucao} '
        else:
            el = OxmlElement('w:t')
            el.text = texto
        run._r.append(el)


class _Conversor:
    def __init__(self, documento):
        from docx.shared import Pt, RGBColor
        self.d = documento
        self.Pt = Pt
        self.RGB = RGBColor

    # ── inline ──────────────────────────────────────────────────────────
    def inline(self, paragrafo, no, fmt=None):
        fmt = dict(fmt or {})
        if isinstance(no, str):
            texto = re.sub(r'\s+', ' ', no)
            if not texto:
                return
            run = paragrafo.add_run(texto)
            run.bold = fmt.get('b') or None
            run.italic = fmt.get('i') or None
            run.underline = fmt.get('u') or None
            run.font.strike = fmt.get('s') or None
            if fmt.get('code'):
                run.font.name = 'Consolas'
            if fmt.get('cor'):
                run.font.color.rgb = self.RGB.from_string(fmt['cor'])
            if fmt.get('tam'):
                run.font.size = self.Pt(fmt['tam'])
            if fmt.get('sup'):
                run.font.superscript = True
            if fmt.get('sub'):
                run.font.subscript = True
            return
        if no.tag == 'br':
            paragrafo.add_run().add_break()
            return
        t = no.tag
        if t == 'strong':
            fmt['b'] = True
        elif t == 'em':
            fmt['i'] = True
        elif t in ('u', 'a'):
            fmt['u'] = True
            if t == 'a':
                fmt['cor'] = '1F5FAD'
        elif t == 's':
            fmt['s'] = True
        elif t == 'code':
            fmt['code'] = True
        elif t in ('sup', 'sub'):
            fmt[t] = True
        elif t == 'span':
            for c in no.classes:
                if c in CORES_SEV:
                    fmt['b'] = True
                    fmt['cor'] = CORES_SEV[c]
                elif c == 'num':
                    fmt['cor'] = COR_PRIMARIA
        for f in no.filhos:
            self.inline(paragrafo, f, fmt)

    def _inlines_sem_strip(self, paragrafo, filhos):
        inicio = len(paragrafo.runs)
        for f in filhos:
            self.inline(paragrafo, f)
        if len(paragrafo.runs) > inicio:
            paragrafo.runs[inicio].text = paragrafo.runs[inicio].text.lstrip()

    def _inlines(self, paragrafo, filhos, fmt=None):
        # remove espaços do começo do parágrafo
        for f in filhos:
            self.inline(paragrafo, f, fmt)
        if paragrafo.runs:
            paragrafo.runs[0].text = paragrafo.runs[0].text.lstrip()

    # ── blocos ──────────────────────────────────────────────────────────
    def blocos(self, no, container=None, nivel_lista=0):
        alvo = container or self.d
        soltos = []

        def descarregar():
            if soltos and any(_texto(x).strip() if not isinstance(x, str) else x.strip() for x in soltos):
                par = alvo.add_paragraph()
                self._inlines(par, list(soltos))
            soltos.clear()

        for f in no.filhos:
            if isinstance(f, str) or f.tag in ('strong', 'em', 'u', 's', 'code', 'span', 'a', 'br', 'sup', 'sub'):
                soltos.append(f)
                continue
            descarregar()
            t = f.tag
            if t in ('h3', 'h4'):
                par = alvo.add_heading(level=2 if t == 'h3' else 3) if alvo is self.d else alvo.add_paragraph()
                self._inlines(par, f.filhos, {'b': alvo is not self.d})
            elif t == 'p':
                par = alvo.add_paragraph()
                self._inlines(par, f.filhos)
            elif t in ('ul', 'ol'):
                self.lista(f, alvo, nivel_lista, numerada=(t == 'ol'))
            elif t == 'table':
                self.tabela(f, alvo)
            elif t == 'div':
                self.callout(f, alvo)
            elif t == 'blockquote':
                par = alvo.add_paragraph(style='Quote') if alvo is self.d else alvo.add_paragraph()
                self._inlines(par, f.filhos, {'i': True})
            elif t == 'pre':
                par = alvo.add_paragraph()
                run = par.add_run(_texto(f))
                run.font.name = 'Consolas'
                run.font.size = self.Pt(8.5)
            else:
                self.blocos(f, alvo, nivel_lista)
        descarregar()

    def lista(self, no, alvo, nivel, numerada):
        # Lista numerada: número escrito no texto. O estilo "List Number" do
        # Word compartilha uma numeração só no documento inteiro — a segunda
        # lista continuaria de onde a primeira parou (11., 27., 33.…).
        from docx.shared import Cm
        estilo = 'List Bullet' + (f' {nivel + 1}' if nivel else '')
        n = 0
        for li in no.filhos:
            if isinstance(li, str) or li.tag != 'li':
                continue
            inline = [x for x in li.filhos if isinstance(x, str) or x.tag not in ('ul', 'ol', 'p', 'table', 'div')]
            if numerada:
                n += 1
                par = alvo.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.9 + 0.6 * nivel)
                par.paragraph_format.first_line_indent = Cm(-0.6)
                par.paragraph_format.space_after = self.Pt(2)
                par.add_run(f'{n}.\t')
                self._inlines_sem_strip(par, inline)
            else:
                try:
                    par = alvo.add_paragraph(style=estilo)
                except KeyError:
                    par = alvo.add_paragraph(style='List Bullet')
                self._inlines(par, inline)
            for sub in li.filhos:
                if not isinstance(sub, str) and sub.tag in ('ul', 'ol'):
                    self.lista(sub, alvo, min(nivel + 1, 2), sub.tag == 'ol')

    def tabela(self, no, alvo):
        from docx.enum.table import WD_TABLE_ALIGNMENT
        linhas = []
        for parte in no.filhos:
            if isinstance(parte, str):
                continue
            if parte.tag == 'tr':
                linhas.append((parte, False))
            else:
                for tr in parte.filhos:
                    if not isinstance(tr, str) and tr.tag == 'tr':
                        linhas.append((tr, parte.tag == 'thead'))
        if not linhas:
            return
        celulas = [[c for c in tr.filhos if not isinstance(c, str) and c.tag in ('td', 'th')] for tr, _ in linhas]
        ncols = max(len(c) for c in celulas) or 1
        tab = alvo.add_table(rows=0, cols=ncols)
        tab.style = 'Table Grid'
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        idx_corpo = 0
        for (tr, cabecalho), cels in zip(linhas, celulas):
            cabecalho = cabecalho or (cels and all(c.tag == 'th' for c in cels))
            row = tab.add_row()
            if cabecalho:
                _repetir_cabecalho(row)
            else:
                idx_corpo += 1
            for i in range(ncols):
                cel = row.cells[i]
                par = cel.paragraphs[0]
                if i < len(cels):
                    self._inlines(par, cels[i].filhos,
                                  {'b': True, 'cor': 'FFFFFF', 'tam': 9} if cabecalho else {'tam': 9})
                if cabecalho:
                    _sombrear(cel, COR_PRIMARIA)
                elif idx_corpo % 2 == 0:
                    _sombrear(cel, COR_ZEBRA)
        self.d.add_paragraph() if alvo is self.d else None

    def callout(self, no, alvo):
        classes = no.classes
        cor = COR_RISCO if 'callout-risco' in classes else COR_CALLOUT
        borda = 'B54708' if 'callout-risco' in classes else COR_PRIMARIA
        tab = alvo.add_table(rows=1, cols=1)
        cel = tab.rows[0].cells[0]
        _borda_esquerda(cel, borda)   # tcBorders precede shd no schema do Word
        _sombrear(cel, cor)
        cel.paragraphs[0]._p.getparent().remove(cel.paragraphs[0]._p)
        self.blocos(no, cel)
        if not cel.paragraphs:
            cel.add_paragraph()
        if alvo is self.d:
            self.d.add_paragraph()


def _estilos(d):
    from docx.enum.style import WD_STYLE_TYPE  # noqa: F401
    from docx.shared import Pt, RGBColor
    normal = d.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    for nome, tam in (('Heading 1', 18), ('Heading 2', 13), ('Heading 3', 11)):
        st = d.styles[nome]
        st.font.name = 'Calibri'
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(COR_PRIMARIA)
        st.paragraph_format.space_before = Pt(14 if nome == 'Heading 1' else 10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def gerar_docx(doc):
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Pt, RGBColor

    ctx = contexto_documento(doc)
    meta = ctx['meta']
    d = Document()
    _estilos(d)
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
    sec.different_first_page_header_footer = True

    # cabeçalho / rodapé (a partir da 2ª página)
    cab = sec.header.paragraphs[0]
    r = cab.add_run(f'{meta.get("empresa", "")} · {meta.get("documento") or doc.titulo}')
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x66, 0x70, 0x85)
    rod = sec.footer.paragraphs[0]
    rod.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = rod.add_run(f'{meta.get("classificacao", "")} · v{doc.versao} · Página ')
    r.font.size = Pt(8)
    r = rod.add_run()
    r.font.size = Pt(8)
    _campo(r, 'PAGE')
    r = rod.add_run(' de ')
    r.font.size = Pt(8)
    r = rod.add_run()
    r.font.size = Pt(8)
    _campo(r, 'NUMPAGES')

    # capa
    for _ in range(3):
        d.add_paragraph()
    par = d.add_paragraph()
    r = par.add_run(meta.get('empresa', ''))
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(COR_PRIMARIA)
    r.bold = True
    par = d.add_paragraph()
    r = par.add_run(meta.get('titulo', doc.titulo))
    r.font.size = Pt(26)
    r.bold = True
    r.font.name = 'Calibri'
    par = d.add_paragraph()
    r = par.add_run(meta.get('subtitulo', ''))
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x47, 0x54, 0x67)
    d.add_paragraph()

    conv = _Conversor(d)
    tab_html = '<table><thead><tr><th>Campo</th><th>Definição</th></tr></thead><tbody>' + ''.join(
        f'<tr><td><strong>{_esc(rotulo)}</strong></td><td>{_esc(valor)}</td></tr>' for rotulo, valor in ctx['campos']
    ) + '</tbody></table>'
    conv.blocos(_arvore(tab_html))
    if meta.get('principio'):
        conv.blocos(_arvore(f'<div class="callout"><p><strong>{_esc(ctx["principio_titulo"])}</strong> '
                            f'{_esc(meta["principio"])}</p></div>'))
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # sumário estático (não depende de o Word atualizar campos)
    d.add_heading('Sumário', level=1)
    for s in ctx['secoes']:
        par = d.add_paragraph()
        par.paragraph_format.space_after = Pt(2)
        r = par.add_run(f'{s["numero"]}. ' if s['numero'] else '')
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(COR_PRIMARIA)
        par.add_run(s['titulo'])
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for s in ctx['secoes']:
        titulo = f'{s["numero"]}. {s["titulo"]}' if s['numero'] else s['titulo']
        d.add_heading(titulo, level=1)
        conv.blocos(_arvore(s['html']))

    buf = io.BytesIO()
    d.core_properties.title = meta.get('documento') or doc.titulo
    d.core_properties.subject = meta.get('titulo', '')
    d.core_properties.author = meta.get('responsavel', '')
    d.save(buf)
    return buf.getvalue()


def _esc(v):
    from django.utils.html import escape
    return escape(v or '—')
