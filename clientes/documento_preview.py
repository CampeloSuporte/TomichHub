"""Pré-visualização dos documentos do cliente (aba Documentos → modal).

O navegador já abre sozinho imagem, PDF, áudio e vídeo — para esses só
devolvemos a URL. O resto é convertido aqui em HTML (sempre escapado) ou em
texto puro, porque o servidor não tem LibreOffice:

- texto/código/config/log/json/xml/csv… → texto (decodificado com fallback)
- markdown → HTML (fonte escapada antes, para não injetar tag)
- .html/.htm → mostrado num iframe sandbox sem script
- .xlsx/.xlsm → tabelas por aba (openpyxl)
- .docx → parágrafos, títulos, listas e tabelas (python-docx)
- .pptx / .odt / .odp / .ods → texto/tabelas lidos direto do XML do zip
- .zip / .tar* → lista de arquivos
- .doc/.xls/.ppt (binários antigos) e o resto → aviso + botão de download
"""
import csv
import io
import os
import tarfile
import zipfile
import xml.etree.ElementTree as ET
from html import escape

LIMITE_TEXTO = 2 * 1024 * 1024      # lê no máximo 2 MB de arquivo texto
LIMITE_LINHAS = 1000                # linhas por planilha/aba
LIMITE_COLUNAS = 60
LIMITE_ARQUIVO_OFFICE = 30 * 1024 * 1024

EXT_IMAGEM = {'png', 'jpg', 'jpeg', 'gif', 'webp', 'bmp', 'svg', 'ico', 'avif', 'jfif'}
EXT_VIDEO = {'mp4', 'webm', 'ogv', 'mov', 'm4v'}
EXT_AUDIO = {'mp3', 'wav', 'ogg', 'oga', 'm4a', 'aac', 'flac', 'opus'}
EXT_HTML = {'html', 'htm'}
EXT_MARKDOWN = {'md', 'markdown'}
EXT_CSV = {'csv', 'tsv'}
EXT_LEGADO = {'doc': 'Word 97-2003', 'xls': 'Excel 97-2003', 'ppt': 'PowerPoint 97-2003'}
EXT_TAR = ('.tar', '.tar.gz', '.tgz', '.tar.bz2', '.tbz2', '.tar.xz', '.txz')


def _ext(nome):
    return os.path.splitext(nome or '')[1].lower().lstrip('.')


def _decodificar(dados):
    enc = 'utf-16' if dados[:2] in (b'\xff\xfe', b'\xfe\xff') else 'utf-8-sig'
    try:
        return dados.decode(enc)
    except UnicodeDecodeError:
        pass
    # Fora do UTF-8, o que aparece aqui é arquivo salvo no Windows em português
    # (cp1252); detector genérico chuta cp1250 e troca "ã" por "ă".
    try:
        return dados.decode('cp1252')
    except UnicodeDecodeError:
        return dados.decode('latin-1', errors='replace')


def _parece_texto(amostra):
    if not amostra:
        return True
    if amostra[:2] in (b'\xff\xfe', b'\xfe\xff'):
        return True
    if b'\x00' in amostra:
        return False
    controles = sum(1 for b in amostra if b < 9 or 13 < b < 32)
    return controles / len(amostra) < 0.05


def _tabela_html(linhas, cabecalho=True):
    if not linhas:
        return '<p class="doc-prev-vazio">(vazia)</p>'
    largura = min(max(len(l) for l in linhas), LIMITE_COLUNAS)
    partes = ['<div class="doc-prev-tabela-wrap"><table class="doc-prev-tabela">']
    for i, linha in enumerate(linhas):
        tag = 'th' if (cabecalho and i == 0) else 'td'
        celulas = list(linha[:largura]) + [''] * (largura - len(linha[:largura]))
        partes.append('<tr>' + ''.join(
            f'<{tag}>{escape("" if c is None else str(c))}</{tag}>' for c in celulas) + '</tr>')
    partes.append('</table></div>')
    return ''.join(partes)


def _abas_html(abas):
    """abas = [(nome, html)] → blocos com título; o JS transforma em abas."""
    return ''.join(
        f'<section class="doc-prev-aba" data-aba="{escape(nome)}">{html}</section>'
        for nome, html in abas
    )


# ---------------------------------------------------------------- formatos

def _xlsx(arquivo):
    import openpyxl
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    abas, cortou = [], False
    for ws in wb.worksheets:
        linhas = []
        for row in ws.iter_rows(values_only=True):
            if len(linhas) >= LIMITE_LINHAS:
                cortou = True
                break
            linhas.append(['' if v is None else v for v in row])
        while linhas and not any(str(v).strip() for v in linhas[-1]):
            linhas.pop()
        if linhas:
            ultima = max((max((i for i, v in enumerate(l) if str(v).strip()), default=-1) for l in linhas))
            linhas = [l[:ultima + 1] for l in linhas]
        abas.append((ws.title, _tabela_html(linhas, cabecalho=False)))
    wb.close()
    aviso = f'Mostrando até {LIMITE_LINHAS} linhas por aba.' if cortou else ''
    return _abas_html(abas), aviso


def _csv(texto, ext):
    amostra = texto[:4096]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters='\t' if ext == 'tsv' else ',;\t|')
    except csv.Error:
        dialeto = csv.excel_tab if ext == 'tsv' else csv.excel
    linhas, cortou = [], False
    for linha in csv.reader(io.StringIO(texto), dialeto):
        if len(linhas) >= LIMITE_LINHAS:
            cortou = True
            break
        linhas.append(linha)
    aviso = f'Mostrando as primeiras {LIMITE_LINHAS} linhas.' if cortou else ''
    return _tabela_html(linhas), aviso


def _docx(arquivo):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(arquivo)

    def runs(p):
        out = []
        for r in p.runs:
            t = escape(r.text)
            if not t:
                continue
            if r.bold:
                t = f'<strong>{t}</strong>'
            if r.italic:
                t = f'<em>{t}</em>'
            if r.underline:
                t = f'<u>{t}</u>'
            out.append(t)
        return ''.join(out) or escape(p.text)

    partes, lista = [], False
    for bloco in d.element.body.iterchildren():
        tag = bloco.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            p = Paragraph(bloco, d)
            estilo = (p.style.name if p.style is not None else '') or ''
            conteudo = runs(p)
            eh_lista = 'List' in estilo or bloco.find('.//{*}numPr') is not None
            if eh_lista and not lista:
                partes.append('<ul>')
                lista = True
            elif not eh_lista and lista:
                partes.append('</ul>')
                lista = False
            if eh_lista:
                partes.append(f'<li>{conteudo}</li>')
            elif estilo.startswith('Heading') or estilo.startswith('Título') or estilo == 'Title':
                nivel = ''.join(ch for ch in estilo if ch.isdigit())[:1] or '1'
                partes.append(f'<h{min(int(nivel) + 2, 6)}>{conteudo}</h{min(int(nivel) + 2, 6)}>')
            elif conteudo.strip():
                partes.append(f'<p>{conteudo}</p>')
            else:
                partes.append('<p>&nbsp;</p>')
        elif tag == 'tbl':
            if lista:
                partes.append('</ul>')
                lista = False
            t = Table(bloco, d)
            linhas = []
            for row in t.rows:
                try:
                    linhas.append([c.text for c in row.cells])
                except Exception:
                    continue
            partes.append(_tabela_html(linhas, cabecalho=False))
    if lista:
        partes.append('</ul>')
    return f'<div class="doc-prev-documento">{"".join(partes)}</div>', ''


def _xml_texto(el):
    return ''.join(el.itertext()).strip()


def _pptx(arquivo):
    with zipfile.ZipFile(arquivo) as z:
        slides = sorted(
            (n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')),
            key=lambda n: int(''.join(ch for ch in os.path.basename(n) if ch.isdigit()) or 0),
        )
        partes = []
        for i, nome in enumerate(slides, 1):
            raiz = ET.fromstring(z.read(nome))
            paragrafos = []
            for p in raiz.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}p'):
                t = _xml_texto(p)
                if t:
                    paragrafos.append(f'<p>{escape(t)}</p>')
            partes.append(f'<div class="doc-prev-slide"><div class="doc-prev-slide-num">Slide {i}</div>'
                          f'{"".join(paragrafos) or "<p class=doc-prev-vazio>(sem texto)</p>"}</div>')
    return ''.join(partes), 'Apresentação mostrada só com o texto de cada slide.'


ODF_TEXT = 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'
ODF_TABLE = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'


def _odf(arquivo, ext):
    with zipfile.ZipFile(arquivo) as z:
        raiz = ET.fromstring(z.read('content.xml'))
    if ext == 'ods':
        abas = []
        for tabela in raiz.iter(f'{{{ODF_TABLE}}}table'):
            linhas = []
            for row in tabela.iter(f'{{{ODF_TABLE}}}table-row'):
                rep_l = int(row.get(f'{{{ODF_TABLE}}}number-rows-repeated', '1'))
                celulas = []
                for c in row:
                    rep = int(c.get(f'{{{ODF_TABLE}}}number-columns-repeated', '1'))
                    celulas.extend([_xml_texto(c)] * min(rep, LIMITE_COLUNAS))
                while celulas and not celulas[-1]:
                    celulas.pop()
                for _ in range(min(rep_l, 50)):
                    linhas.append(celulas)
                if len(linhas) >= LIMITE_LINHAS:
                    break
            while linhas and not linhas[-1]:
                linhas.pop()
            abas.append((tabela.get(f'{{{ODF_TABLE}}}name', 'Planilha'), _tabela_html(linhas, cabecalho=False)))
        return _abas_html(abas), ''
    partes = []
    for el in raiz.iter():
        if el.tag in (f'{{{ODF_TEXT}}}h', f'{{{ODF_TEXT}}}p'):
            t = _xml_texto(el)
            if el.tag.endswith('}h'):
                partes.append(f'<h4>{escape(t)}</h4>')
            else:
                partes.append(f'<p>{escape(t) or "&nbsp;"}</p>')
    aviso = 'Apresentação mostrada só com o texto.' if ext == 'odp' else ''
    return f'<div class="doc-prev-documento">{"".join(partes)}</div>', aviso


def _rtf(texto):
    import re
    texto = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes([int(m.group(1), 16)]).decode('cp1252', 'replace'), texto)
    texto = re.sub(r'\\par[d]?\b ?', '\n', texto)
    texto = re.sub(r'\{\\\*[^{}]*\}', '', texto)
    texto = re.sub(r'\\[a-zA-Z]+-?\d* ?', '', texto)
    texto = texto.replace('{', '').replace('}', '')
    return texto.strip()


def _lista_arquivos(itens):
    linhas = [['Arquivo', 'Tamanho']]
    for nome, tam in itens[:LIMITE_LINHAS]:
        linhas.append([nome, '' if tam is None else _tamanho(tam)])
    aviso = f'Mostrando {LIMITE_LINHAS} de {len(itens)} itens.' if len(itens) > LIMITE_LINHAS else ''
    return _tabela_html(linhas), aviso


def _tamanho(n):
    for un in ('B', 'KB', 'MB', 'GB'):
        if n < 1024:
            return f'{n:.0f} {un}' if un == 'B' else f'{n:.1f} {un}'
        n /= 1024
    return f'{n:.1f} TB'


# ---------------------------------------------------------------- entrada

def montar_preview(documento):
    """Devolve o dict que o modal consome (ver _renderizarPreviewDocumento no listar.html)."""
    f = documento.arquivo
    nome = documento.nome or os.path.basename(f.name)
    ext = _ext(nome) or _ext(f.name)
    base = {'ok': True, 'nome': nome, 'ext': ext, 'url': f.url, 'aviso': ''}

    try:
        tamanho = f.size
    except (OSError, FileNotFoundError):
        return {'ok': False, 'erro': 'Arquivo não encontrado no servidor.'}
    base['tamanho'] = _tamanho(tamanho)

    if ext in EXT_IMAGEM:
        return {**base, 'tipo': 'imagem'}
    if ext == 'pdf':
        return {**base, 'tipo': 'pdf'}
    if ext in EXT_VIDEO:
        return {**base, 'tipo': 'video'}
    if ext in EXT_AUDIO:
        return {**base, 'tipo': 'audio'}
    if ext in EXT_LEGADO:
        return {**base, 'tipo': 'nao_suportado',
                'erro': f'Formato {EXT_LEGADO[ext]} (.{ext}) não tem pré-visualização. '
                        f'Salve como .{ext}x para abrir aqui, ou baixe o arquivo.'}

    escritorio = {'xlsx', 'xlsm', 'docx', 'pptx', 'odt', 'ods', 'odp', 'zip'}
    if (ext in escritorio or nome.lower().endswith(EXT_TAR)) and tamanho > LIMITE_ARQUIVO_OFFICE:
        return {**base, 'tipo': 'nao_suportado', 'erro': 'Arquivo grande demais para pré-visualizar.'}

    try:
        with f.open('rb') as fh:
            if ext in ('xlsx', 'xlsm'):
                html, aviso = _xlsx(fh)
                return {**base, 'tipo': 'html', 'html': html, 'aviso': aviso}
            if ext == 'docx':
                html, aviso = _docx(fh)
                return {**base, 'tipo': 'html', 'html': html, 'aviso': aviso}
            if ext == 'pptx':
                html, aviso = _pptx(fh)
                return {**base, 'tipo': 'html', 'html': html, 'aviso': aviso}
            if ext in ('odt', 'ods', 'odp'):
                html, aviso = _odf(fh, ext)
                return {**base, 'tipo': 'html', 'html': html, 'aviso': aviso}
            if ext == 'zip':
                with zipfile.ZipFile(fh) as z:
                    itens = [(i.filename, None if i.is_dir() else i.file_size) for i in z.infolist()]
                html, aviso = _lista_arquivos(itens)
                return {**base, 'tipo': 'html', 'html': html, 'aviso': aviso}
            if nome.lower().endswith(EXT_TAR):
                with tarfile.open(fileobj=fh, mode='r:*') as t:
                    itens = [(m.name, m.size if m.isfile() else None) for m in t.getmembers()]
                html, aviso = _lista_arquivos(itens)
                return {**base, 'tipo': 'html', 'html': html, 'aviso': aviso}

            dados = fh.read(LIMITE_TEXTO + 1)
    except Exception as e:  # arquivo corrompido / formato diferente da extensão
        return {**base, 'tipo': 'nao_suportado', 'erro': f'Não foi possível ler o arquivo ({type(e).__name__}).'}

    cortado = len(dados) > LIMITE_TEXTO
    dados = dados[:LIMITE_TEXTO]
    aviso = 'Arquivo grande: mostrando só os primeiros 2 MB.' if cortado else ''

    if not _parece_texto(dados[:8192]):
        return {**base, 'tipo': 'nao_suportado',
                'erro': 'Este tipo de arquivo não tem pré-visualização. Use o botão de download.'}

    texto = _decodificar(dados)
    if ext in EXT_HTML:
        return {**base, 'tipo': 'iframe_html', 'texto': texto, 'aviso': aviso}
    if ext in EXT_MARKDOWN:
        import markdown
        html = markdown.markdown(escape(texto, quote=False), extensions=['tables', 'fenced_code'])
        return {**base, 'tipo': 'html', 'html': f'<div class="doc-prev-documento">{html}</div>', 'aviso': aviso}
    if ext in EXT_CSV:
        html, aviso_csv = _csv(texto, ext)
        return {**base, 'tipo': 'html', 'html': html, 'aviso': aviso or aviso_csv}
    if ext == 'rtf':
        return {**base, 'tipo': 'texto', 'texto': _rtf(texto), 'aviso': aviso}
    return {**base, 'tipo': 'texto', 'texto': texto, 'aviso': aviso}
